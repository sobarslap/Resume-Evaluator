"""Extract plain text from resume files held in memory (bytes)."""

from __future__ import annotations

import io

from docx import Document
from pypdf import PdfReader

from .config import SUPPORTED_SUFFIXES


class UnsupportedFileError(ValueError):
    """Raised when a file type we cannot read is supplied."""


def read_pdf(data: bytes) -> str:
    """Extract text from PDF bytes."""
    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            parts.append(page_text)
    return "\n".join(parts)


def read_docx(data: bytes) -> str:
    """Extract text from DOCX bytes, including table cells."""
    document = Document(io.BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def _suffix(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot != -1 else ""


def extract_text(filename: str, data: bytes) -> str:
    """Dispatch to the right reader based on the file extension.

    Raises UnsupportedFileError for anything that is not a PDF or DOCX.
    """
    suffix = _suffix(filename)
    if suffix == ".pdf":
        return read_pdf(data)
    if suffix == ".docx":
        return read_docx(data)
    raise UnsupportedFileError(
        f"Unsupported file type '{suffix or filename}'. "
        f"Supported types: {', '.join(SUPPORTED_SUFFIXES)}."
    )
