"""Tests for text extraction (no network needed)."""

import io

import pytest
from docx import Document

from evaluator.extractors import UnsupportedFileError, extract_text, read_docx


def _make_docx(lines: list[str]) -> bytes:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_read_docx_extracts_paragraphs():
    data = _make_docx(["Jane Doe", "Skills: Python, AWS"])
    text = read_docx(data)
    assert "Jane Doe" in text
    assert "Python, AWS" in text


def test_read_docx_includes_table_cells():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Experience"
    table.rows[0].cells[1].text = "3 years"
    buf = io.BytesIO()
    doc.save(buf)
    text = read_docx(buf.getvalue())
    assert "Experience" in text
    assert "3 years" in text


def test_extract_text_dispatches_docx():
    data = _make_docx(["Hello"])
    assert "Hello" in extract_text("resume.docx", data)


def test_extract_text_rejects_unsupported():
    with pytest.raises(UnsupportedFileError):
        extract_text("resume.txt", b"plain text")


def test_extract_text_is_case_insensitive():
    data = _make_docx(["Caps"])
    assert "Caps" in extract_text("RESUME.DOCX", data)
